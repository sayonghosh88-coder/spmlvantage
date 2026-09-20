"""Excel adapter and auditable portfolio calculations for the nine-sheet tracker."""
from __future__ import annotations
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
import json
import math
import zipfile
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

SCHEMA = json.loads((Path(__file__).parent / 'schema.json').read_text(encoding='utf-8'))
REGISTRY = 'Project Status & EVM'
CHILDREN = [s['name'] for s in SCHEMA if s['name'] != REGISTRY]
SCHEMAS = {s['name']: s for s in SCHEMA}
RANK = {'red': 0, 'amber': 1, 'green': 2, 'unknown': 3}
LABEL = {'red': 'Critical', 'amber': 'At risk', 'green': 'On track', 'unknown': 'Not confirmed'}

class DataError(ValueError):
    pass

def number(value):
    if value is None or value == '' or isinstance(value, bool):
        return None
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None

def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            pass
    return None

def evm(p):
    names = ['BAC approved cost budget (INR lakh)', 'PV planned value (INR lakh)', 'EV earned value (INR lakh)', 'AC actual cost (INR lakh)']
    b, pv, ev, ac = [number(p.get(k)) for k in names]
    valid = all(v is not None and v >= 0 for v in (b, pv, ev, ac)) and b > 0 and pv <= b and ev <= b
    if not valid:
        return dict(BAC=b, PV=pv, EV=ev, AC=ac, CPI=None, SPI=None, CV=None, SV=None, EAC=None, ETC=None, VAC=None, valid=False)
    cpi, spi = (ev / ac if ac else None), (ev / pv if pv else None)
    eac = b / cpi if cpi else None
    return dict(BAC=b, PV=pv, EV=ev, AC=ac, CPI=cpi, SPI=spi, CV=ev-ac, SV=ev-pv,
                EAC=eac, ETC=eac-ac if eac is not None else None, VAC=b-eac if eac is not None else None, valid=True)

def recalculate(rows, sheet):
    for r in rows:
        if sheet == REGISTRY:
            e = evm(r)
            for field, key in {'CPI cost performance index':'CPI','SPI schedule performance index':'SPI',
                'Cost variance (INR lakh)':'CV','Schedule variance (INR lakh)':'SV','EAC forecast final cost (INR lakh)':'EAC',
                'ETC forecast remaining cost (INR lakh)':'ETC','VAC budget less forecast cost (INR lakh)':'VAC'}.items():
                r[field] = e[key]
        elif sheet == 'Construction & Resources':
            plan, actual = number(r.get('Weekly target quantity')), number(r.get('Actual quantity achieved'))
            r['Shortfall quantity'] = max(plan-actual,0) if plan is not None and actual is not None else None
        elif sheet == 'Procurement Status':
            need, received = number(r.get('Total required quantity')), number(r.get('Quantity received at site'))
            r['Balance quantity to receive'] = need-received if need is not None and received is not None else None

def read_workbook(payload):
    if len(payload) > 20*1024*1024:
        raise DataError('Use an Excel workbook smaller than 20 MB.')
    try:
        with zipfile.ZipFile(BytesIO(payload)) as z:
            if sum(i.file_size for i in z.infolist()) > 100*1024*1024 or len(z.infolist()) > 10000:
                raise DataError('The expanded workbook is too large.')
        wb = load_workbook(BytesIO(payload), data_only=False, keep_links=False)
    except DataError:
        raise
    except Exception as exc:
        raise DataError('This file could not be read as an .xlsx workbook.') from exc
    data, warnings = {}, []
    try:
        for spec in SCHEMA:
            name = spec['name']
            if name not in wb.sheetnames:
                raise DataError(f'Missing sheet: {name}. Please use the redesigned nine-sheet template.')
            ws = wb[name]
            if ws.max_row > 20000 or ws.max_column > 200:
                raise DataError(f'{name}: limit the input to 20,000 rows and 200 columns.')
            header = next((n for n in range(1, min(40,ws.max_row)+1)
                           if ws.cell(n,1).value == 'Project code' and all(h in [c.value for c in ws[n]] for h in spec['headers'])),None)
            if not header:
                raise DataError(f'{name}: the expected column headings were not found. Retain the template headings.')
            headings = [c.value for c in ws[header]]
            if any(headings.count(h)>1 for h in spec['headers']):
                raise DataError(f'{name}: duplicate column headings.')
            end = spec['end']
            for table in ws.tables.values():
                _, top, _, bottom = range_boundaries(table.ref)
                if top == header:
                    end = bottom
                    break
            fields = {h:headings.index(h)+1 for h in spec['headers']}
            records = []
            for n in range(header+1,end+1):
                cells = {h:ws.cell(n,c) for h,c in fields.items()}
                values = {h:c.value for h,c in cells.items() if c.data_type != 'f'}
                if not any(v is not None and v != '' for v in values.values()):
                    continue
                if not values.get('Project code'):
                    raise DataError(f'{name}, row {n}: enter a project code for this record.')
                for h,c in cells.items():
                    if c.data_type == 'f' and h not in spec['formulas']:
                        raise DataError(f'{name}, {c.coordinate}: enter a verified value instead of an input formula.')
                r = {h:values.get(h) for h in spec['headers']}
                r['Project code'] = str(r['Project code']).strip()
                r['_row'] = n
                records.append(r)
            recalculate(records,name)
            data[name] = records
        codes = [r['Project code'] for r in data[REGISTRY]]
        if not codes or len(codes) != len(set(codes)):
            raise DataError('Project Status & EVM needs exactly one current row per project code.')
        keys = {'MD Decisions':['Decision ID'],'Construction & Resources':['Front / activity ID','Week starting'],
                'Engineering Status':['Drawing / document reference'],'Procurement Status':['Procurement package / item ID'],
                'Invoicing Status':['Invoice / RA bill reference'],'Liability Status':['Liability / bill reference'],
                'Lookahead Plan':['Front / activity ID','Week starting'],'Daily MOM':['Item / action ID']}
        for name in CHILDREN:
            seen = set()
            for r in data[name]:
                if r['Project code'] not in codes:
                    raise DataError(f'{name}, row {r["_row"]}: unknown project code {r["Project code"]}.')
                parts = [r.get(f) for f in keys[name]]
                if any(x is None or x == '' for x in parts):
                    warnings.append(f'{name}, row {r["_row"]}: reference or period is missing; record retained.')
                    continue
                key = (r['Project code'], *parts)
                if key in seen:
                    raise DataError(f'{name}, row {r["_row"]}: duplicate record key. Update the existing record.')
                seen.add(key)
        snapshots=[]
        ws=wb['Construction & Resources']
        table=ws.tables.get('ProjectDeployment')
        if table:
            left, top, right, bottom=range_boundaries(table.ref)
            heads=[ws.cell(top,c).value for c in range(left,right+1)]
            for n in range(top+1,bottom+1):
                row={h:ws.cell(n,left+i).value for i,h in enumerate(heads)}
                if row.get('Project code'):
                    if row['Project code'] not in codes:
                        raise DataError(f'Deployment snapshot row {n}: unknown project code.')
                    row['_row']=n
                    snapshots.append(row)
        dates=[as_date(r.get('Last updated')) for rs in data.values() for r in rs]
        latest=max((d for d in dates if d),default=date.today())
        fictional=any('FICTIONAL TEST DATA' in str(r.get('Source / entry notes','')).upper() for rs in data.values() for r in rs)
        return {'sheets':data,'snapshots':snapshots,'latest':latest,'fictional':fictional,'warnings':warnings}
    finally:
        wb.close()

def rows(data,sheet,code=None):
    rs=data['sheets'].get(sheet,[])
    return [r for r in rs if code is None or r['Project code']==code]

def total(rs,field):
    values=[number(r.get(field)) for r in rs]
    known=[v for v in values if v is not None]
    return {'value':sum(known) if known else None,'known':len(known),'count':len(values)}

def progress_rag(plan,actual):
    plan,actual=number(plan),number(actual)
    if plan is None or actual is None or not (0<=plan<=1 and 0<=actual<=1):
        return 'unknown'
    gap=round((actual-plan)*100,8)
    return 'green' if gap>=-3 else 'amber' if gap>=-8 else 'red'

def index_rag(value):
    return 'unknown' if value is None else 'green' if value>=1 else 'amber' if value>=.9 else 'red'

def slip_days(p):
    approved,forecast=as_date(p.get('Latest approved completion')),as_date(p.get('Forecast completion'))
    return (forecast-approved).days if approved and forecast else None

def project_rag(p):
    e=evm(p)
    slip=slip_days(p)
    signals=[progress_rag(p.get('Overall planned progress'),p.get('Overall actual progress')),index_rag(e['CPI']),index_rag(e['SPI'])]
    if slip is not None:
        signals.append('red' if slip>30 else 'amber' if slip>0 else 'green')
    known=[s for s in signals if s!='unknown']
    return min(known,key=RANK.get) if known else 'unknown'

def ordered_projects(data):
    return sorted(rows(data,REGISTRY),key=lambda p:(RANK[project_rag(p)],
        (number(p.get('Overall actual progress')) or 0)-(number(p.get('Overall planned progress')) or 0),p['Project code']))

def portfolio_evm(projects):
    valid=[p for p in projects if evm(p)['valid']]
    dates={as_date(p.get('EVM reporting date')) for p in valid}
    if len(valid)!=len(projects) or len(dates)!=1 or None in dates:
        return {'CPI':None,'SPI':None,'note':'Complete comparable EVM inputs and align reporting dates.','count':len(valid)}
    vals=[evm(p) for p in valid]
    earned=sum(e['EV'] for e in vals);cost=sum(e['AC'] for e in vals);planned=sum(e['PV'] for e in vals)
    return {'CPI':earned/cost if cost else None,'SPI':earned/planned if planned else None,
            'note':'ΣEV / ΣPV for SPI; ΣEV / ΣAC for CPI.','count':len(valid)}

def decision_queue(data):
    projects={p['Project code']:p for p in rows(data,REGISTRY)}
    return sorted([r for r in rows(data,'MD Decisions') if r.get('Status')=='Awaiting decision'],
                  key=lambda r:(RANK[project_rag(projects[r['Project code']])],as_date(r.get('Result due date')) or date.max,r['Project code'],r.get('Decision ID','')))

def record_state(sheet,r,review):
    status=r.get({'MD Decisions':'Status','Construction & Resources':'Front readiness','Lookahead Plan':'Activity status'}.get(sheet,'Present Status')) or 'Not confirmed'
    completed=status in ['Closed','Completed','Approved','Fully received','Fully paid','Fully collected','Not applicable']
    date_field={'MD Decisions':'Result due date','Construction & Resources':'Committed execution date','Engineering Status':'Required-for-construction date',
                'Procurement Status':'Required-at-site date','Invoicing Status':'Contractual payment due date','Liability Status':'Contractual payment due date',
                'Lookahead Plan':'Planned finish','Daily MOM':'Target closure date'}[sheet]
    due=as_date(r.get(date_field))
    if sheet=='Construction & Resources':
        target,actual=number(r.get('Weekly target quantity')),number(r.get('Actual quantity achieved'))
        if target is not None and target>0 and actual is not None and actual>=target:
            return 'green','Target achieved',due
    if completed:
        return 'green',status,due
    if status in ['Blocked','On hold','Disputed'] or (due and due<review):
        return 'red',status,due
    if sheet=='Construction & Resources':
        if any(r.get(k)=='NO' for k in ['Drawing Approval Status','Contractor Deployment Status','Shuttering Availability Status','Steel Availability Status']):
            return 'red',status,due
        return ('green' if status=='Ready' else 'amber'),status,due
    if sheet=='Lookahead Plan':
        clearance,start=as_date(r.get('Constraint clearance target date')),as_date(r.get('Planned start'))
        if r.get('Execution readiness')=='Blocked' or (clearance and start and clearance>start):
            return 'red',status,due
        return ('green' if r.get('Execution readiness')=='Ready' else 'amber'),status,due
    return 'amber',status,due

def ordered_records(data,sheet,code,review):
    return sorted(rows(data,sheet,code),key=lambda r:(RANK[record_state(sheet,r,review)[0]],record_state(sheet,r,review)[2] or date.max,r['_row']))

def deployment(data,code):
    rs=[r for r in data['snapshots'] if r.get('Project code')==code]
    latest=max((as_date(r.get('Reporting date')) for r in rs if as_date(r.get('Reporting date'))),default=None)
    current=[r for r in rs if as_date(r.get('Reporting date'))==latest] if latest else []
    return current,latest

MACHINE_FIELDS=['Batching plant','Ajax','Excavator','Miller','Truck / dumper','Tractor / breaker','JCB','Other machines']

def deployment_totals(data,code=None):
    codes=[code] if code else [p['Project code'] for p in rows(data,REGISTRY)]
    current=[];dates=[]
    for project_code in codes:
        snapshot,d=deployment(data,project_code)
        current.extend(snapshot)
        if d: dates.append(d)
    machines=[number(r.get(f)) for r in current for f in MACHINE_FIELDS]
    return {'planned':total(current,'Labour planned')['value'],'actual':total(current,'Labour actual')['value'],
            'machines':sum(machines) if machines and all(n is not None for n in machines) else None,
            'sites':len(current),'projects':len(dates),'expected_projects':len(codes),'dates':set(dates)}
